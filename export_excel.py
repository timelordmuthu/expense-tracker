"""
export_excel.py  —  Multi-sheet Excel export.
Hidden bonuses are excluded from every sheet.
Open bonuses appear in the Income sheet and Overview.
"""

import os
import math
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from data_manager import load_data

OUTPUT_PATH = "expense_report.xlsx"

# ── Palette ───────────────────────────────────────────────────────────────────
C_NAVY    = "1A1A2E"
C_NAVY2   = "2D2D4E"
C_GREEN   = "1E6B3C"
C_RED     = "8B1A1A"
C_BLUE    = "1A3A5C"
C_AMBER   = "7A4F1A"
C_CREAM   = "E8E4D9"
C_ALT     = "F7F6F2"
C_WHITE   = "FFFFFF"
C_BLACK   = "000000"
C_LINK    = "008000"
C_ORANGE  = "FFF3E0"

# ── Style helpers ─────────────────────────────────────────────────────────────

def _f(bold=False, size=10, color=C_BLACK, name="Arial", italic=False):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)

def _fill(c):
    return PatternFill("solid", start_color=c, fgColor=c)

def _border(s="thin"):
    side = Side(style=s, color="D0CCC0")
    return Border(left=side, right=side, top=side, bottom=side)

def _ctr():  return Alignment(horizontal="center", vertical="center", wrap_text=True)
def _lft():  return Alignment(horizontal="left",   vertical="center", wrap_text=True)
def _rgt():  return Alignment(horizontal="right",  vertical="center")

INR = '₹#,##0.00;(₹#,##0.00);"-"'

def _hdr_row(ws, row, labels, bg, fg=C_WHITE):
    for c, lbl in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=lbl)
        cell.font = _f(bold=True, color=fg); cell.fill = _fill(bg)
        cell.alignment = _ctr(); cell.border = _border()

def _data_row(ws, row, vals, alt=False):
    bg = C_ALT if alt else C_WHITE
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.fill = _fill(bg); cell.border = _border()
        if isinstance(v, float):
            cell.number_format = INR; cell.alignment = _rgt()
            cell.font = _f(name="Courier New", size=9)
        else:
            cell.alignment = _lft(); cell.font = _f(size=9)

def _total_row(ws, row, vals):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = _f(bold=True, name="Courier New", size=9)
        cell.fill = _fill(C_ORANGE); cell.border = _border("medium")
        cell.alignment = _rgt() if c > 1 else _lft()
        if isinstance(v, str) and v.startswith("="):
            cell.number_format = INR

def _title(ws, title, sub, bg, cols=8):
    ws.merge_cells(f"A1:{get_column_letter(cols)}1")
    ws.merge_cells(f"A2:{get_column_letter(cols)}2")
    t = ws["A1"]; t.value = title
    t.font = _f(bold=True, size=14, color=C_WHITE)
    t.fill = _fill(bg); t.alignment = _ctr()
    s = ws["A2"]; s.value = sub
    s.font = _f(italic=True, size=9, color=C_CREAM)
    s.fill = _fill(C_NAVY2); s.alignment = _ctr()
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 15

def _safe(v):
    try: return round(float(v or 0), 2)
    except: return 0.0

def _dt(iso):
    if not iso: return "—"
    try: return datetime.fromisoformat(iso).strftime("%d %b %Y")
    except: return str(iso)[:10]

# ── Sheet builders ────────────────────────────────────────────────────────────

def _sheet_expenses(wb, data):
    ws = wb.create_sheet("Expenses")
    _title(ws, "💸 Expenses Ledger",
           f"All recorded outflows  ·  Generated {datetime.now().strftime('%d %b %Y %H:%M')}",
           C_RED)
    _hdr_row(ws, 3, ["#","Date","Category","Note","Visibility","Status","Amount (₹)"], C_RED)

    entries = sorted(data.get("expenses", []), key=lambda x: x.get("date",""))
    ds = 4
    for i, e in enumerate(entries, 1):
        _data_row(ws, ds+i-1, [
            i,
            _dt(e.get("date") or e.get("created_at")),
            e.get("category","Other"),
            e.get("note") or "—",
            "Hidden" if e.get("hidden") else "Open",
            (e.get("status") or "").capitalize(),
            _safe(e.get("amount")),
        ], alt=i%2==0)

    total_row = ds + len(entries)
    last      = total_row - 1
    gc        = get_column_letter(7)
    _total_row(ws, total_row, ["TOTAL","","","","","",
        f"=SUM({gc}{ds}:{gc}{last})" if entries else 0.0])

    for col, w in {"A":5,"B":14,"C":20,"D":30,"E":10,"F":12,"G":15}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:G{total_row}"


def _sheet_income(wb, data):
    ws = wb.create_sheet("Income")
    _title(ws, "💵 Income & Open Bonuses",
           f"Inflows only — hidden bonuses excluded  ·  {datetime.now().strftime('%d %b %Y %H:%M')}",
           C_GREEN)
    _hdr_row(ws, 3, ["#","Date","Source","Frequency","Note","Status","Amount (₹)"], C_GREEN)

    inc   = [dict(e, _type="Income") for e in data.get("income", [])]
    bonus = [dict(e, _type="Bonus")  for e in data.get("bonuses",[]) if not e.get("hidden")]
    rows  = sorted(inc + bonus, key=lambda x: x.get("date") or x.get("created_at",""))

    ds = 4
    for i, e in enumerate(rows, 1):
        _data_row(ws, ds+i-1, [
            i,
            _dt(e.get("date") or e.get("created_at")),
            e.get("source") or e.get("_type","Income"),
            e.get("mode","—"),
            e.get("note") or "—",
            (e.get("status") or "").capitalize(),
            _safe(e.get("amount")),
        ], alt=i%2==0)

    total_row = ds + len(rows)
    last      = total_row - 1
    gc        = get_column_letter(7)
    _total_row(ws, total_row, ["TOTAL","","","","","",
        f"=SUM({gc}{ds}:{gc}{last})" if rows else 0.0])

    for col, w in {"A":5,"B":14,"C":18,"D":12,"E":28,"F":12,"G":15}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:G{total_row}"


def _sheet_summary(wb, data):
    ws = wb.create_sheet("Category Summary")
    _title(ws, "📊 Category-wise Summary",
           f"Monthly & yearly spend breakdown  ·  {datetime.now().strftime('%d %b %Y %H:%M')}",
           C_BLUE, cols=3)

    expenses = data.get("expenses", [])
    rows = []
    for e in expenses:
        try: dt = datetime.fromisoformat(e.get("date") or e.get("created_at",""))
        except: dt = None
        rows.append({
            "category": e.get("category") or "Other",
            "amount":   _safe(e.get("amount")),
            "month":    dt.strftime("%b %Y") if dt else "Unknown",
            "year":     dt.year if dt else None,
            "hidden":   bool(e.get("hidden")),
        })

    df = pd.DataFrame(rows) if rows else pd.DataFrame(
        columns=["category","amount","month","year","hidden"])

    # ── Monthly ──
    ws.merge_cells("A4:C4")
    c = ws["A4"]; c.value = "Monthly Breakdown by Category"
    c.font = _f(bold=True, size=11, color=C_WHITE)
    c.fill = _fill(C_BLUE); c.alignment = _ctr()
    ws.row_dimensions[4].height = 18

    _hdr_row(ws, 5, ["Category", "Month", "Total (₹)"], C_NAVY2)

    if not df.empty:
        monthly = (df.groupby(["category","month"], sort=False)["amount"]
                   .sum().reset_index().sort_values(["month","category"]))
        ms = 6
        for i, (_, r) in enumerate(monthly.iterrows(), 1):
            _data_row(ws, ms+i-1, [r["category"], r["month"], float(r["amount"])], alt=i%2==0)
        me = ms + len(monthly) - 1
        _total_row(ws, me+1, ["TOTAL","", f"=SUM(C{ms}:C{me})"])

        # ── Yearly ──
        yr = me + 4
        ws.merge_cells(f"A{yr}:C{yr}")
        c2 = ws[f"A{yr}"]; c2.value = "Yearly Breakdown by Category"
        c2.font = _f(bold=True, size=11, color=C_WHITE)
        c2.fill = _fill(C_BLUE); c2.alignment = _ctr()
        ws.row_dimensions[yr].height = 18

        _hdr_row(ws, yr+1, ["Category","Year","Total (₹)"], C_NAVY2)

        yearly = (df[df["year"].notna()].groupby(["category","year"])["amount"]
                  .sum().reset_index().sort_values(["year","category"]))
        ys = yr + 2
        for i, (_, r) in enumerate(yearly.iterrows(), 1):
            _data_row(ws, ys+i-1, [r["category"], int(r["year"]), float(r["amount"])], alt=i%2==0)
        ye = ys + len(yearly) - 1
        _total_row(ws, ye+1, ["TOTAL","", f"=SUM(C{ys}:C{ye})"])
    else:
        ws.cell(row=6, column=1, value="No expense data yet.").font = _f(italic=True)

    for col, w in {"A":24,"B":14,"C":16}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A6"


def _sheet_overview(wb, data):
    ws = wb.active
    ws.title = "Overview"
    _title(ws, "📋 Expense Tracker — Report Overview",
           f"Generated {datetime.now().strftime('%d %b %Y %H:%M')}", C_NAVY, cols=2)

    open_bonus_total = sum(
        _safe(b.get("amount")) for b in data.get("bonuses",[]) if not b.get("hidden")
    )
    hidden_count = sum(1 for b in data.get("bonuses",[]) if b.get("hidden"))
    inc_rows  = len(data.get("income",[]))
    exp_rows  = len(data.get("expenses",[]))
    bonus_vis = sum(1 for b in data.get("bonuses",[]) if not b.get("hidden"))

    kv = [
        ("Report Date",        datetime.now().strftime("%d %b %Y")),
        ("Total Income",       f"=Income!G{inc_rows+bonus_vis+4}"),
        ("Total Expenses",     f"=Expenses!G{exp_rows+4}"),
        ("Net Balance",        "=Overview!B5-Overview!B6"),
        ("Open Bonus Total",   open_bonus_total),
        ("Hidden Bonuses",     f"{hidden_count} (amounts not disclosed)"),
        ("Day Streak 🔥",      data.get("streak",0)),
        ("Points ⭐",          data.get("points",0)),
        ("Categories",         len(data.get("categories",[]))),
        ("Rewards Tracked",    len(data.get("rewards",[]))),
    ]

    for i, (lbl, val) in enumerate(kv, 4):
        alt = i % 2 == 0
        lc = ws.cell(row=i, column=1, value=lbl)
        vc = ws.cell(row=i, column=2, value=val)
        for cell in (lc, vc):
            cell.fill = _fill(C_ALT if alt else C_WHITE)
            cell.border = _border()
        lc.font = _f(bold=True); lc.alignment = _lft()
        vc.alignment = _rgt()
        is_formula = isinstance(val, str) and val.startswith("=")
        vc.font = _f(name="Courier New", color=C_LINK if is_formula else C_BLACK)
        if lbl in ("Total Income","Total Expenses","Net Balance","Open Bonus Total"):
            vc.number_format = INR

    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 22

# ── Public API ────────────────────────────────────────────────────────────────

def export_to_excel(output_path: str = OUTPUT_PATH) -> str:
    """
    Build and save the 4-sheet Excel report.
    Returns the absolute path of the written file.
    Hidden bonuses never appear in any sheet.
    """
    data = load_data()
    wb   = Workbook()

    _sheet_overview(wb, data)
    _sheet_expenses(wb, data)
    _sheet_income(wb, data)
    _sheet_summary(wb, data)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)) or ".", exist_ok=True)
    wb.save(output_path)
    return os.path.abspath(output_path)


if __name__ == "__main__":
    p = export_to_excel()
    print(f"✓ Report saved → {p}")
